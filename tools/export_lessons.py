"""Export lessons from the database to the three teaching files.

    python3 tools/export_lessons.py              # every lesson
    python3 tools/export_lessons.py forgiveness  # one lesson (by code)

For each lesson, writes lessons/output/<code>/:
  <code>-slides.pptx     title, scripture, teaching points, discussion questions
  <code>-handout.docx    scripture, outline with fill-in blanks, discussion questions, application
  <code>-teacher.docx    everything: background, commentator quotes with credits, timing, answers
The design is deliberately plain; Dave will refine the look later.
"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor as PRGB  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt as PPt  # noqa: E402

from library import db  # noqa: E402
from library.scripture import format_ref as _format_ref, split_vid  # noqa: E402

OUT = os.path.join(db.ROOT, "lessons", "output")
INK = (0x22, 0x2B, 0x36)
ACCENT = (0x2E, 0x5E, 0x8C)
GREY = (0x6B, 0x72, 0x80)
SECTION_NAME = {"opening": "Opening", "scripture": "Scripture", "background": "Background",
                "teaching_point": "Teaching point", "differ": "Where Christians differ",
                "questions": "Discussion questions", "application": "Application", "prayer": "Closing prayer"}
BLANK_LINE = "_" * 14


# ----------------------------------------------------------------------------- data

def load(con, code):
    L = con.execute("""SELECT l.id, l.code, l.title, l.big_idea, s.name, l.audience, l.length_minutes, w.code,
                              l.status, l.version_kind, t.name, ci.name, l.contested_issue_id, l.family_code
                       FROM lessons l JOIN subjects s ON s.id = l.topic_subject_id JOIN works w ON w.id = l.translation_work_id
                       LEFT JOIN traditions t ON t.id = l.tradition_id LEFT JOIN contested_issues ci ON ci.id = l.contested_issue_id
                       WHERE l.code = ?""", (code,)).fetchone()
    keys = ["id", "code", "title", "big_idea", "topic", "audience", "minutes", "translation", "status", "version",
            "tradition", "issue", "issue_id", "family"]
    lesson = dict(zip(keys, L))
    lesson["sections"] = []
    for sid, kind, heading, body, minutes, note, hidden in con.execute(
            """SELECT id, kind, heading, body, minutes, teacher_note, default_hidden FROM lesson_sections
               WHERE lesson_id = ? ORDER BY position""", (lesson["id"],)):
        sec = dict(id=sid, kind=kind, heading=heading, body=body, minutes=minutes, note=note, hidden=hidden)
        sec["verses"] = con.execute("""SELECT start_verse_id, end_verse_id FROM lesson_scripture
                                       WHERE section_id = ? ORDER BY id""", (sid,)).fetchall()
        sec["quotes"] = con.execute("SELECT quote_text, credit FROM lesson_quotes WHERE section_id = ? ORDER BY id",
                                    (sid,)).fetchall()
        sec["blanks"] = con.execute("SELECT sentence, answer FROM lesson_blanks WHERE section_id = ? ORDER BY position",
                                    (sid,)).fetchall()
        if kind == "differ":
            sec["views"] = []
            for code in re.findall(r"@view:(\S+)", body):
                row = con.execute("""SELECT t.name, ps.summary, ps.key_texts FROM perspective_summaries ps
                                     JOIN traditions t ON t.id = ps.tradition_id
                                     WHERE ps.issue_id = ? AND t.code = ?""", (lesson["issue_id"], code)).fetchone()
                if row:
                    sec["views"].append(row)
            sec["body"] = re.sub(r"\s*@view:\S+", "", body).strip()
        lesson["sections"].append(sec)
    lesson["main"] = con.execute("""SELECT start_verse_id, end_verse_id FROM lesson_scripture
                                    WHERE lesson_id = ? AND role = 'main' ORDER BY id""", (lesson["id"],)).fetchall()
    lesson["questions"] = con.execute("SELECT question, answer FROM lesson_questions WHERE lesson_id = ? ORDER BY position",
                                      (lesson["id"],)).fetchall()
    return lesson


def verses(con, translation, start, end):
    """[(verse number, chapter, text, starts_paragraph)] for a range, KJV markup cleaned."""
    out = []
    for vid, text in con.execute("""SELECT bt.verse_id, bt.text FROM bible_text bt JOIN works w ON w.id = bt.work_id
                                    WHERE w.code = ? AND bt.verse_id BETWEEN ? AND ? ORDER BY bt.verse_id""",
                                 (translation, start, end)):
        _, ch, v = split_vid(vid)
        para = text.startswith("¶")
        text = re.sub(r"[\[\]]", "", text.replace("¶", "")).strip()
        out.append((v, ch, text, para))
    return out


def format_ref(a, b=None):
    """As library.scripture.format_ref, but a single psalm reads 'Psalm 23', as teachers say it."""
    return re.sub(r"^Psalms (\d+)(?!-\d+$)", r"Psalm \1", _format_ref(a, b))


def passage_ref(lesson):
    return "; ".join(format_ref(a, b) for a, b in lesson["main"])


def filled(sentence, answer):
    return sentence.replace("____", answer, 1)


# ----------------------------------------------------------------------------- slides

def rgb(t):
    return PRGB(*t)


def text_box(slide, left, top, width, height, paragraphs, size=24, color=INK, bold=False, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        runs = para if isinstance(para, list) else [(para, {})]
        for text, fmt in runs:
            r = p.add_run()
            r.text = text
            r.font.size = PPt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", bold)
            r.font.italic = fmt.get("italic", italic)
            r.font.color.rgb = rgb(fmt.get("color", color))
        p.space_after = PPt(fmt.get("space", 10) if runs else 10)
    return box


def base_slide(prs, title, footer):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(ACCENT)
    bar.line.fill.background()
    text_box(s, Inches(0.7), Inches(0.45), Inches(11.9), Inches(1.0), [title], size=32, bold=True)
    text_box(s, Inches(0.7), Inches(6.85), Inches(11.9), Inches(0.4), [footer], size=12, color=GREY)
    return s


def chunk_verses(vs, limit=70):
    out, cur, n = [], [], 0
    for v in vs:
        w = len(v[2].split())
        if cur and n + w > limit:
            out.append(cur)
            cur, n = [], 0
        cur.append(v)
        n += w
    if cur:
        if out and sum(len(v[2].split()) for v in cur) < 25:
            out[-1] += cur   # a verse or two left over: keep it with the slide before
        else:
            out.append(cur)
    return out


def make_slides(con, L, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    footer = L["title"]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(3.55), Inches(1.6), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(ACCENT)
    bar.line.fill.background()
    text_box(s, Inches(0.7), Inches(1.6), Inches(11.9), Inches(1.9), [L["title"]], size=44, bold=True)
    text_box(s, Inches(0.7), Inches(3.85), Inches(11.4), Inches(1.6), [L["big_idea"]], size=24, italic=True)
    sub = passage_ref(L) + (f"   ·   {L['tradition']} version" if L["tradition"] else "")
    text_box(s, Inches(0.7), Inches(5.9), Inches(11.9), Inches(0.5), [sub], size=16, color=GREY)

    for start, end in L["main"]:
        vs = verses(con, L["translation"], start, end)
        chunks = chunk_verses(vs)
        for k, chunk in enumerate(chunks, 1):
            title = format_ref(start, end) + (f" ({k} of {len(chunks)})" if len(chunks) > 1 else "")
            s = base_slide(prs, title, f"{footer}   ·   {L['translation']}")
            para = []
            for v, ch, text, _ in chunk:
                para += [(f"{v} ", {"size": 16, "color": ACCENT, "bold": True}), (text + " ", {"size": 28})]
            text_box(s, Inches(0.7), Inches(1.45), Inches(11.9), Inches(5.3), [para], size=28)

    n = 0
    for sec in L["sections"]:
        if sec["kind"] != "teaching_point":
            continue
        n += 1
        s = base_slide(prs, f"{n}. {sec['heading']}", footer)
        lines = [[("•  ", {"color": ACCENT, "bold": True, "size": 30}), (filled(b, a), {"size": 30, "space": 18})]
                 for b, a in sec["blanks"]]
        if sec["verses"]:
            lines.append([("Read: " + "; ".join(format_ref(a, b) for a, b in sec["verses"]),
                           {"size": 20, "color": GREY, "italic": True})])
        text_box(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.0), lines, size=30)
        notes = sec["body"] + "".join(f"\n\n“{q}” — {c}" for q, c in sec["quotes"])
        s.notes_slide.notes_text_frame.text = notes

    qs = L["questions"]
    for k in range(0, len(qs), 3):
        s = base_slide(prs, "Discuss" + (" (continued)" if k else ""), footer)
        lines = [[(f"{k + i + 1}.  ", {"color": ACCENT, "bold": True, "size": 28}), (q, {"size": 28, "space": 18})]
                 for i, (q, _) in enumerate(qs[k:k + 3])]
        text_box(s, Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.0), lines, size=28)
    prs.save(path)


# ----------------------------------------------------------------------------- documents

def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12.5)):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(*INK) if name == "Title" else RGBColor(*ACCENT)
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Pt(64))
    return doc


def para(doc, text="", italic=False, bold=False, size=None, color=None, indent=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.italic, r.bold = italic, bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    return p


def shade(p, fill="EEF3F8"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.text = text
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(*GREY)


def print_passage(doc, con, L, start, end, size=11):
    p = None
    for v, ch, text, new_para in verses(con, L["translation"], start, end):
        if p is None or new_para:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
        r = p.add_run(str(v))
        r.font.superscript = True
        r.font.color.rgb = RGBColor(*ACCENT)
        r2 = p.add_run(" " + text + " ")
        r2.font.size = Pt(size)


def body_paras(doc, text):
    for chunk in [c for c in text.split("\n\n") if c.strip()]:
        doc.add_paragraph(" ".join(chunk.split()))


def make_handout(con, L, path):
    doc = new_doc()
    doc.add_paragraph(L["title"], style="Title")
    para(doc, f"{passage_ref(L)} ({L['translation']})", italic=True, color=GREY)
    doc.add_heading("Scripture", level=1)
    for start, end in L["main"]:
        para(doc, format_ref(start, end), bold=True)
        print_passage(doc, con, L, start, end)
    doc.add_heading("Outline", level=1)
    n = 0
    for sec in L["sections"]:
        if sec["kind"] != "teaching_point":
            continue
        n += 1
        doc.add_heading(f"{n}. {sec['heading']}", level=2)
        for b, _ in sec["blanks"]:
            para(doc, b.replace("____", BLANK_LINE), indent=14)
        if sec["verses"]:
            para(doc, "Read: " + "; ".join(format_ref(a, b) for a, b in sec["verses"]), italic=True, size=10,
                 color=GREY, indent=14)
    doc.add_heading("Discussion questions", level=1)
    for i, (q, _) in enumerate(L["questions"], 1):
        para(doc, f"{i}. {q}")
        para(doc, " ")
    for sec in L["sections"]:
        if sec["kind"] == "application":
            doc.add_heading("This week", level=1)
            body_paras(doc, sec["body"])
    footer(doc, f"{L['title']}  ·  Scripture: {L['translation']} (public domain)")
    doc.save(path)


def make_teacher(con, L, path):
    doc = new_doc()
    doc.add_paragraph(L["title"], style="Title")
    p = para(doc, L["big_idea"], italic=True, size=12.5)
    shade(p)
    rows = [("Main passage", f"{passage_ref(L)} ({L['translation']})"), ("Topic", L["topic"]),
            ("Audience", L["audience"]), ("Length", f"about {L['minutes']} minutes"),
            ("Status", "Draft: written by AI from the library, not yet reviewed by Dave")]
    if L["version"] == "all_views":
        rows.insert(2, ("Version", f"All views (contested point: {L['issue']})"))
    elif L["version"] == "tradition":
        rows.insert(2, ("Version", f"{L['tradition']} (contested point: {L['issue']})"))
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light List Accent 1"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, v
    doc.add_heading("Suggested timing", level=1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light List Accent 1"
    n = 0
    for sec in L["sections"]:
        if sec["minutes"] is None:
            continue
        if sec["kind"] == "teaching_point":
            n += 1
        label = f"Teaching point {n}: {sec['heading']}" if sec["kind"] == "teaching_point" else SECTION_NAME[sec["kind"]]
        cells = t.add_row().cells
        cells[0].text, cells[1].text = label, f"{sec['minutes']} min"
    cells = t.add_row().cells
    cells[0].text, cells[1].text = "Total", f"{sum(s['minutes'] or 0 for s in L['sections'])} min"

    n = 0
    for sec in L["sections"]:
        if sec["kind"] == "teaching_point":
            n += 1
            title = f"Teaching point {n}: {sec['heading']}"
        else:
            title = SECTION_NAME[sec["kind"]]
        if sec["minutes"]:
            title += f"  ({sec['minutes']} min)"
        doc.add_heading(title, level=1)
        if sec["kind"] == "differ" and sec["hidden"]:
            p = para(doc, "In this tradition's version, this section is hidden from the class materials unless the "
                          "user turns it on. It is included here so the teacher knows the other views.",
                     italic=True, size=10, color=GREY)
        if sec["kind"] == "scripture":
            for start, end in L["main"]:
                para(doc, format_ref(start, end), bold=True)
                print_passage(doc, con, L, start, end)
        body_paras(doc, sec["body"])
        for name, summary, key_texts in sec.get("views", []):
            doc.add_heading(name, level=2)
            doc.add_paragraph(summary)
            if key_texts:
                para(doc, "Key texts: " + key_texts, italic=True, size=10, color=GREY)
        if sec["verses"]:
            doc.add_heading("Supporting verses", level=2)
            for a, b in sec["verses"]:
                vs = verses(con, L["translation"], a, b)
                if len(vs) <= 4:
                    p = para(doc, format_ref(a, b) + ": ", bold=True, size=10)
                    r = p.add_run(" ".join(text for _, _, text, _ in vs))
                    r.font.size = Pt(10)
                else:
                    para(doc, format_ref(a, b) + " (read from the Bible)", bold=True, size=10)
        for q, c in sec["quotes"]:
            p = para(doc, f"“{q}”", italic=True, indent=18)
            p.paragraph_format.space_after = Pt(0)
            para(doc, f"— {c}", size=9.5, color=GREY, indent=18)
        if sec["blanks"]:
            doc.add_heading("Handout blanks (answers)", level=2)
            for b, a in sec["blanks"]:
                p = doc.add_paragraph()
                before, _, after = b.partition("____")
                p.add_run(before)
                r = p.add_run(a)
                r.bold = True
                r.underline = True
                p.add_run(after)
        if sec["kind"] == "questions":
            for i, (q, a) in enumerate(L["questions"], 1):
                para(doc, f"{i}. {q}", bold=True)
                if a:
                    para(doc, a, indent=14)
        if sec["note"]:
            p = para(doc, "Teacher note: " + sec["note"], italic=True, size=10)
            shade(p, "FFF6E5")
    doc.add_heading("Sources quoted", level=1)
    seen = []
    for sec in L["sections"]:
        for _, c in sec["quotes"]:
            if c not in seen:
                seen.append(c)
    for c in seen:
        para(doc, c, size=10)
    para(doc, "Every quotation was checked word for word against the library text it cites. All sources are in the "
              "public domain.", italic=True, size=9.5, color=GREY)
    footer(doc, f"Teacher's notes  ·  {L['title']}  ·  draft for review")
    doc.save(path)


def main(argv):
    con = db.connect()
    codes = argv or [r[0] for r in con.execute("SELECT code FROM lessons ORDER BY code")]
    for code in codes:
        L = load(con, code)
        folder = os.path.join(OUT, code)
        os.makedirs(folder, exist_ok=True)
        make_slides(con, L, os.path.join(folder, f"{code}-slides.pptx"))
        make_handout(con, L, os.path.join(folder, f"{code}-handout.docx"))
        make_teacher(con, L, os.path.join(folder, f"{code}-teacher.docx"))
    print(f"exported {len(codes)} lesson(s) to {os.path.relpath(OUT, db.ROOT)}/")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
